"""Fase 4: Excel acumulativo con formato (una hoja por etapa del pipeline)."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ESTILO_HOJAS = {
    "Desviaciones": "BDD7EE",
    "Recepcion_Cruce": "BDD7EE",
    "Categorizacion": "BDD7EE",
    "Asignacion": "BDD7EE",
    "Tabla_Final": "000000",
}


def _ordenar_por_fecha(df, columna="Fin"):
    if df is None or df.empty or columna not in df.columns:
        return df
    orden = pd.to_datetime(df[columna], format="%d/%m/%Y", errors="coerce")
    return df.assign(_o=orden).sort_values("_o", kind="stable", na_position="first").drop(columns="_o").reset_index(drop=True)


def _acumular(ruta: Path, hoja: str, df_nuevo):
    """Concatena las filas nuevas con lo que ya tenía esa hoja en el archivo."""
    previo = None
    if ruta.exists():
        try:
            previo = pd.read_excel(ruta, sheet_name=hoja, dtype=str).fillna("")
        except ValueError:
            previo = None
    if df_nuevo is not None:
        df_nuevo = df_nuevo.astype(str).replace("nan", "")
    partes = [d for d in (previo, df_nuevo) if d is not None and not d.empty]
    # drop_duplicates: volver a correr la misma fecha no duplica filas
    return pd.concat(partes, ignore_index=True).drop_duplicates(ignore_index=True) if partes else None


def _formatear(ws, color_encabezado: str):
    texto_blanco = color_encabezado == "000000"
    for i, columna in enumerate(ws.columns, start=1):
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        ws.column_dimensions[get_column_letter(i)].width = min(largo + 3, 60)
    for celda in ws[1]:
        celda.fill = PatternFill(fill_type="solid", fgColor=color_encabezado)
        celda.font = Font(bold=True, color="FFFFFF" if texto_blanco else "000000")
        celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def guardar_reporte(ruta, hojas: dict) -> Path:
    """Guarda (acumulando con corridas anteriores) cada dataframe en su hoja.

    `hojas` es un dict {nombre_hoja: dataframe | None}.
    """
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)

    finales = {}
    for nombre, df in hojas.items():
        acumulado = _ordenar_por_fecha(_acumular(ruta, nombre, df), "FIN" if nombre == "Tabla_Final" else "Fin")
        if acumulado is not None:
            finales[nombre] = acumulado

    if not finales:
        raise ValueError("No hay datos para guardar en el reporte.")

    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        for nombre, df in finales.items():
            df.to_excel(writer, sheet_name=nombre, index=False)
        libro = writer.book
        for nombre in finales:
            _formatear(libro[nombre], ESTILO_HOJAS.get(nombre, "BDD7EE"))
        if "Tabla_Final" in finales:
            libro.active = libro.sheetnames.index("Tabla_Final")
    return ruta
